#!/usr/bin/env python3
"""
DWH -> Destination reconciliation tool.

Compares source tables (DWH) with destination tables (other databases) without
copying data between them. Each database summarises its own data; only the
summaries are compared, and rows are pulled only for the small slices that differ.

  Level 1  Row counts per partition
  Level 2  Aggregate fingerprints per partition (sums, null counts, text lengths,
           min/max dates, distinct keys)
  Level 3  Drill-down: mismatched partitions are split into key buckets, bucket
           fingerprints are compared, and only mismatched buckets are fetched and
           diffed row by row. A few random buckets are also spot-checked row by row
           every run, even when everything matches.

Usage:
  python recon.py -c recon_config.yaml
  python recon.py -c recon_config.yaml --table policy_master --level 2
  python recon.py -c recon_config.yaml --show-sql          # print SQL, run nothing

Exit codes: 0 = all matched, 1 = mismatches found, 2 = one or more tables errored.
"""
from __future__ import annotations

import argparse
import concurrent.futures as cf
import datetime as dt
import html
import logging
import math
import os
import random
import re
import sqlite3
import sys
import time
import urllib.parse
from dataclasses import dataclass, field
from decimal import Decimal

import pandas as pd
import yaml

log = logging.getLogger("recon")


# =============================================================================
# SQL dialects: every engine-specific expression lives here
# =============================================================================
class Dialect:
    name = "generic"

    def mod(self, a, n):            return f"MOD({a}, {n})"
    def trim(self, x):              return f"TRIM({x})"
    def length(self, x):            return f"LENGTH({x})"
    def ascii(self, x):             return f"ASCII({x})"
    def substr(self, x, start, n):  return f"SUBSTR({x}, {start}, {n})"
    def concat(self, a, b):         return f"({a} || {b})"
    def right(self, x, n):          return f"RIGHT({x}, {n})"
    def to_day(self, x):            return f"CAST({x} AS DATE)"
    def fmt_day(self, x):           return f"TO_CHAR({x}, 'YYYY-MM-DD')"
    def fmt_month(self, x):         return f"TO_CHAR({x}, 'YYYY-MM')"
    def date_lit(self, s):          return f"CAST('{s}' AS TIMESTAMP)"
    def num(self, x, scale):        return f"ROUND({x}, {scale})"
    def bigint(self, x):            return x
    def upper(self, x):             return f"UPPER({x})"
    def epoch_days(self, x):        return f"(CAST({x} AS DATE) - DATE '1970-01-01')"
    def epoch_secs(self, x):        return f"CAST(EXTRACT(EPOCH FROM {x}) AS BIGINT)"

    # Shared building blocks -------------------------------------------------
    def text_len(self, x):
        """Length of trimmed text; empty string treated as NULL (Oracle-compatible)."""
        return self.bigint(self.length(f"NULLIF({self.trim(x)}, '')"))

    def text_sig(self, x, ignore_case=False):
        """Cheap content signature: codes of 1st, 2nd and last character. Catches value
        changes that keep the same length (e.g. LAPSED -> PAIDUP)."""
        t = f"NULLIF({self.trim(x)}, '')"
        if ignore_case:
            t = self.upper(t)
        return self.bigint(f"({self.ascii(t)} + 7 * COALESCE({self.ascii(self.substr(t, 2, 1))}, 0)"
                           f" + 49 * {self.ascii(self.right(t, 1))})")

    def bucket_expr(self, key, key_type, n):
        if key_type == "number":
            return self.mod(key, n)
        # Text key: mix the codes of the last 4 characters. Padding with '####'
        # makes short keys behave identically on every engine.
        last4 = self.right(self.concat("'####'", self.trim(key)), 4)
        parts = [f"{w} * {self.ascii(self.substr(last4, i, 1))}"
                 for i, w in zip(range(1, 5), (1, 7, 49, 343))]
        return self.mod("(" + " + ".join(parts) + ")", n)


class OracleDialect(Dialect):
    name = "oracle"
    def right(self, x, n):          return f"SUBSTR({x}, -{n})"
    def to_day(self, x):            return f"TRUNC({x})"
    def date_lit(self, s):          return f"TO_DATE('{s}', 'YYYY-MM-DD')"
    def epoch_days(self, x):        return f"(TRUNC(CAST({x} AS DATE)) - DATE '1970-01-01')"
    def epoch_secs(self, x):        return f"ROUND((CAST({x} AS DATE) - DATE '1970-01-01') * 86400)"


class SqlServerDialect(Dialect):
    name = "mssql"
    def mod(self, a, n):            return f"(({a}) % {n})"
    def trim(self, x):              return f"LTRIM(RTRIM({x}))"
    def length(self, x):            return f"LEN({x})"
    def substr(self, x, start, n):  return f"SUBSTRING({x}, {start}, {n})"
    def concat(self, a, b):         return f"({a} + {b})"
    def fmt_day(self, x):           return f"CONVERT(CHAR(10), {x}, 23)"
    def fmt_month(self, x):         return f"CONVERT(CHAR(7), {x}, 23)"
    def date_lit(self, s):          return f"CAST('{s}' AS DATETIME2)"
    def num(self, x, scale):        return f"CAST(ROUND({x}, {scale}) AS DECIMAL(38, {scale}))"
    def bigint(self, x):            return f"CAST({x} AS BIGINT)"
    def ascii(self, x):             return f"UNICODE({x})"
    def epoch_days(self, x):        return f"CAST(DATEDIFF(DAY, '1970-01-01', {x}) AS BIGINT)"
    def epoch_secs(self, x):        return f"DATEDIFF_BIG(SECOND, '1970-01-01', {x})"


class PostgresDialect(Dialect):
    name = "postgresql"
    def num(self, x, scale):        return f"ROUND(CAST({x} AS NUMERIC), {scale})"


class MySqlDialect(Dialect):
    name = "mysql"
    def length(self, x):            return f"CHAR_LENGTH({x})"
    def substr(self, x, start, n):  return f"SUBSTRING({x}, {start}, {n})"
    def concat(self, a, b):         return f"CONCAT({a}, {b})"
    def to_day(self, x):            return f"DATE({x})"
    def fmt_day(self, x):           return f"DATE_FORMAT({x}, '%Y-%m-%d')"
    def fmt_month(self, x):         return f"DATE_FORMAT({x}, '%Y-%m')"
    def date_lit(self, s):          return f"'{s}'"
    def epoch_days(self, x):        return f"DATEDIFF({x}, '1970-01-01')"
    def epoch_secs(self, x):        return f"TIMESTAMPDIFF(SECOND, '1970-01-01', {x})"


class SqliteDialect(Dialect):
    name = "sqlite"
    def mod(self, a, n):            return f"(({a}) % {n})"
    def ascii(self, x):             return f"UNICODE({x})"
    def right(self, x, n):          return f"SUBSTR({x}, -{n})"
    def to_day(self, x):            return f"DATE({x})"
    def fmt_day(self, x):           return f"STRFTIME('%Y-%m-%d', {x})"
    def fmt_month(self, x):         return f"STRFTIME('%Y-%m', {x})"
    def date_lit(self, s):          return f"'{s}'"
    def epoch_days(self, x):        return f"CAST(JULIANDAY(DATE({x})) - 2440587.5 AS INTEGER)"
    def epoch_secs(self, x):        return f"CAST(ROUND((JULIANDAY({x}) - 2440587.5) * 86400) AS INTEGER)"


DIALECTS = {
    "oracle": OracleDialect(), "mssql": SqlServerDialect(), "sqlserver": SqlServerDialect(),
    "postgresql": PostgresDialect(), "postgres": PostgresDialect(),
    "mysql": MySqlDialect(), "mariadb": MySqlDialect(), "sqlite": SqliteDialect(),
}


# =============================================================================
# Database connections
# =============================================================================
class Database:
    """Thin read-only wrapper: SQLAlchemy for real servers, stdlib sqlite3 for files."""

    def __init__(self, name: str, cfg: dict):
        self.name = name
        self.type = str(cfg["type"]).lower()
        if self.type not in DIALECTS:
            raise ValueError(f"Connection '{name}': unsupported type '{self.type}'. "
                             f"Use one of: {', '.join(sorted(DIALECTS))}")
        self.dialect = DIALECTS[self.type]
        self.engine = None
        self.sqlite_path = None
        if self.type == "sqlite":
            self.sqlite_path = cfg["path"]
            if not os.path.exists(self.sqlite_path):
                raise FileNotFoundError(f"Connection '{name}': {self.sqlite_path} not found")
            return
        try:
            from sqlalchemy import create_engine
        except ImportError as e:
            raise ImportError("SQLAlchemy is required for server databases: "
                              "pip install sqlalchemy") from e
        if self.type == "oracle":
            try:  # bigger fetch batches = far fewer network round trips
                import oracledb
                oracledb.defaults.arraysize = 10000
                oracledb.defaults.prefetchrows = 10000
                if cfg.get("thick_mode") and oracledb.is_thin_mode():
                    # Needs Oracle Instant Client; only for DBs that reject thin mode
                    oracledb.init_oracle_client(lib_dir=cfg.get("client_lib_dir") or None)
            except ImportError:
                pass
        kw = {"fast_executemany": True} if self.type in ("mssql", "sqlserver") else {}
        self.engine = create_engine(self._url(cfg), pool_pre_ping=True,
                                    pool_size=int(cfg.get("pool_size", 8)), max_overflow=4, **kw)

    def _password(self, cfg, required=True):
        env = cfg.get("password_env")
        pwd = os.environ.get(env) if env else cfg.get("password")
        if pwd is None and required:
            raise ValueError(f"Connection '{self.name}': password not found. Add a line "
                             f"'{env or 'PASSWORD_VAR'}=yourpassword' to the .env file next to "
                             f"the config (or set it as an environment variable)")
        if pwd is not None and is_placeholder(pwd):
            raise ValueError(f"Connection '{self.name}': the password for {env} in .env is "
                             f"still the placeholder '{pwd}'. Put the real password there.")
        return pwd

    def _url(self, cfg):
        for k, v in cfg.items():
            if k != "password" and is_placeholder(v):
                raise ValueError(f"Connection '{self.name}': '{k}' in recon_config.yaml is "
                                 f"still '{v}'. Fill in the real value (see the comment next to it).")
        if "url" in cfg:                                   # full SQLAlchemy url given
            url = cfg["url"]
            if "{password}" in url:
                url = url.replace("{password}", urllib.parse.quote_plus(self._password(cfg)))
            return url
        from sqlalchemy.engine import URL                  # build from separate fields
        t = self.type
        if t == "oracle":
            q = {"service_name": str(cfg["service_name"])} if cfg.get("service_name") else {}
            return URL.create("oracle+oracledb", username=cfg["user"], password=self._password(cfg),
                              host=cfg["host"], port=int(cfg.get("port", 1521)),
                              database=None if q else cfg.get("sid"), query=q)
        if t in ("mssql", "sqlserver"):
            windows = bool(cfg.get("windows_auth", False))
            q = {"driver": str(cfg.get("driver", "ODBC Driver 18 for SQL Server")),
                 "TrustServerCertificate": "yes"}
            if windows:
                q["Trusted_Connection"] = "yes"
            port = cfg.get("port")
            return URL.create("mssql+pyodbc",
                              username=None if windows else cfg["user"],
                              password=None if windows else self._password(cfg),
                              host=str(cfg["server"]), port=int(port) if port else None,
                              database=cfg["database"], query=q)
        if t in ("postgresql", "postgres"):
            return URL.create("postgresql+psycopg2", username=cfg["user"],
                              password=self._password(cfg), host=cfg["host"],
                              port=int(cfg.get("port", 5432)), database=cfg["database"])
        if t in ("mysql", "mariadb"):
            return URL.create("mysql+pymysql", username=cfg["user"], password=self._password(cfg),
                              host=cfg["host"], port=int(cfg.get("port", 3306)),
                              database=cfg["database"])
        raise ValueError(f"Connection '{self.name}': give either 'url' or the connection fields")

    def query(self, sql: str) -> list[tuple]:
        log.debug("[%s] %s", self.name, sql)
        if self.sqlite_path:
            con = sqlite3.connect(f"file:{self.sqlite_path}?mode=ro", uri=True)
            try:
                return con.execute(sql).fetchall()
            finally:
                con.close()
        with self.engine.connect() as c:
            return [tuple(r) for r in c.exec_driver_sql(sql).fetchall()]

    def write_tables(self, frames: dict, schema: str | None = None):
        """Append several DataFrames in ONE transaction: a run is saved completely or not at all."""
        frames = {t: df for t, df in frames.items() if not df.empty}
        if self.sqlite_path:
            con = sqlite3.connect(self.sqlite_path)
            try:
                with con:
                    for t, df in frames.items():
                        df.to_sql(t, con, if_exists="append", index=False)
            finally:
                con.close()
        else:
            with self.engine.begin() as c:
                for t, df in frames.items():
                    df.to_sql(t, c, schema=schema, if_exists="append", index=False, chunksize=1000)


# =============================================================================
# Table specification from config
# =============================================================================
@dataclass
class Column:
    name: str            # logical name (used in reports)
    src: str
    tgt: str
    type: str            # number | text | date | datetime
    scale: int = 2
    tolerance: float = 0.0
    ignore_case: bool = False
    signature: bool = True


@dataclass
class Side:
    db: Database
    table: str
    filter: str | None
    keys: list[str]
    part_col: str | None

    @property
    def d(self) -> Dialect:
        return self.db.dialect


@dataclass
class TableSpec:
    name: str
    src: Side
    tgt: Side
    key_names: list[str]
    key_type: str
    columns: list[Column]
    grain: str | None    # day | month | value | None


def _pairs(obj):
    """keys / columns can be a list of names or a {source_name: target_name} mapping."""
    if isinstance(obj, list):
        return [(k, k) if isinstance(k, str) else next(iter(k.items())) for k in obj]
    return list(obj.items())


def is_placeholder(v) -> bool:
    return isinstance(v, str) and v.strip().lower().strip("x") == "" and len(v.strip()) >= 3


def _check_placeholders(obj, where):
    if isinstance(obj, dict):
        for k, v in obj.items():
            if is_placeholder(k) or is_placeholder(v):
                raise ValueError(f"{where}: '{k}' is still a placeholder ({v!r}). "
                                 f"Fill in the real value in recon_config.yaml.")
            _check_placeholders(v, where)
    elif isinstance(obj, list):
        for v in obj:
            if is_placeholder(v):
                raise ValueError(f"{where}: still contains placeholder {v!r}. "
                                 f"Fill in the real value in recon_config.yaml.")
            _check_placeholders(v, where)


_IDENT = r'(?:[A-Za-z_][A-Za-z0-9_$#]*|"[^"]+"|\[[^\]]+\])'
IDENT_RE = re.compile(rf"^{_IDENT}(?:\.{_IDENT}){{0,2}}$")


def _check_ident(value, what, where):
    if not isinstance(value, str) or not IDENT_RE.match(value.strip()):
        raise ValueError(f"{where}: {what} {value!r} is not a valid table/column name. Use "
                         f"letters, digits and _ (schema.table allowed), or \"quoted\" / "
                         f"[bracketed] names.")


def _check_filter(value, where):
    if value and (";" in value or "--" in value or "/*" in value):
        raise ValueError(f"{where}: filter {value!r} may not contain ';', '--' or '/*'. "
                         f"A filter must be a single condition like  ISSUE_DATE >= '2025-04-01'.")


def build_spec(t: dict, dbs: dict, settings: dict) -> TableSpec:
    name = t["name"]
    where = f"Table '{name}'"
    _check_placeholders(t, where)
    for side_key in ("source", "target"):
        side_cfg = t.get(side_key) or {}
        _check_ident(side_cfg.get("table"), f"{side_key} table", where)
        _check_filter(side_cfg.get("filter"), where)
    _check_filter(t.get("filter"), where)
    for k in _pairs(t["keys"]):
        for v in k:
            _check_ident(v, "key column", where)
    for cname, cdef in (t.get("columns") or {}).items():
        _check_ident(cname, "column", where)
        if isinstance(cdef, dict) and cdef.get("target"):
            _check_ident(cdef["target"], "target column", where)
    part = t.get("partition") or {}
    for k in ("column", "target_column"):
        if part.get(k):
            _check_ident(part[k], "partition column", where)
    s_cfg, t_cfg = t["source"], t["target"]
    for side in (s_cfg, t_cfg):
        if side["connection"] not in dbs:
            raise ValueError(f"Table '{name}': unknown connection '{side['connection']}'")

    keys = _pairs(t["keys"])
    cols = []
    default_scale = int(settings.get("number_scale", 2))
    for cname, cdef in (t.get("columns") or {}).items():
        if isinstance(cdef, str):
            cdef = {"type": cdef}
        ctype = cdef.get("type", "text").lower()
        if ctype not in ("number", "text", "date", "datetime"):
            raise ValueError(f"Table '{name}', column '{cname}': type must be "
                             f"number/text/date/datetime")
        cols.append(Column(name=cname, src=cname, tgt=cdef.get("target", cname), type=ctype,
                           scale=int(cdef.get("scale", default_scale)),
                           tolerance=float(cdef.get("tolerance", 0)),
                           ignore_case=bool(cdef.get("ignore_case", False)),
                           signature=bool(cdef.get("signature",
                                                   settings.get("text_signature", True)))))

    part = t.get("partition") or {}
    grain = part.get("grain", "day").lower() if part else None
    if grain not in (None, "day", "month", "value"):
        raise ValueError(f"Table '{name}': partition grain must be day/month/value")
    sp = part.get("column")
    tp = part.get("target_column", sp)

    common = t.get("filter")
    def side(cfg, key_idx, pcol):
        f = [x for x in (common, cfg.get("filter")) if x]
        return Side(db=dbs[cfg["connection"]], table=cfg["table"],
                    filter=" AND ".join(f"({x})" for x in f) or None,
                    keys=[k[key_idx] for k in keys], part_col=pcol)

    return TableSpec(name=name, src=side(s_cfg, 0, sp), tgt=side(t_cfg, 1, tp),
                     key_names=[k[0] for k in keys],
                     key_type=str(t.get("key_type", "text")).lower(),
                     columns=cols, grain=grain if sp else None)


# =============================================================================
# SQL builders
# =============================================================================
def _lit(v):
    if v is None:
        return "NULL"
    if isinstance(v, (int, float, Decimal)) and not isinstance(v, bool):
        return str(v)
    return "'" + str(v).replace("'", "''") + "'"


def _where(side: Side, extra: list[str] | None = None) -> str:
    conds = [c for c in [side.filter] + (extra or []) if c]
    return (" WHERE " + " AND ".join(conds)) if conds else ""


def _part_expr(side: Side, grain):
    d, c = side.d, side.part_col
    return {"day": d.fmt_day(c), "month": d.fmt_month(c), "value": c}[grain]


def partition_filter(side: Side, grain, pkey, praw) -> str | None:
    if grain is None:
        return None
    c, d = side.part_col, side.d
    if pkey is None:
        return f"{c} IS NULL"
    if grain == "value":
        return f"{c} = {_lit(praw)}"
    if grain == "day":
        start = dt.date.fromisoformat(pkey)
        end = start + dt.timedelta(days=1)
    else:  # month
        y, m = map(int, pkey.split("-"))
        start = dt.date(y, m, 1)
        end = dt.date(y + (m == 12), m % 12 + 1, 1)
    return f"{c} >= {d.date_lit(start.isoformat())} AND {c} < {d.date_lit(end.isoformat())}"


def _metric_exprs(side: Side, spec: TableSpec, level: int, is_src: bool):
    """(inner select expressions, outer aggregate expressions, metric names)."""
    d = side.d
    inner, outer, names = [], ["COUNT(*)"], ["rows"]
    if level < 2:
        return inner, outer, names
    if len(side.keys) == 1:
        inner.append(f"{side.keys[0]} AS k0")
        outer.append("COUNT(DISTINCT k0)")
        names.append("distinct_keys")
    for i, col in enumerate(spec.columns):
        c = col.src if is_src else col.tgt
        a = f"m{i}"
        if col.type == "number":
            inner.append(f"{d.num(c, col.scale)} AS {a}")
            outer += [f"SUM({a})", f"COUNT({a})"]
            names += [f"sum({col.name})", f"non_null({col.name})"]
        elif col.type == "text":
            inner.append(f"{d.text_len(c)} AS {a}")
            outer += [f"SUM({a})", f"COUNT({a})"]
            names += [f"text_len({col.name})", f"non_blank({col.name})"]
            if col.signature:
                inner.append(f"{d.text_sig(c, col.ignore_case)} AS {a}s")
                outer.append(f"SUM({a}s)")
                names.append(f"text_sig({col.name})")
        else:
            ep = d.epoch_days(c) if col.type == "date" else d.epoch_secs(c)
            inner += [f"{d.to_day(c) if col.type == 'date' else c} AS {a}", f"{ep} AS {a}e"]
            outer += [f"MIN({a})", f"MAX({a})", f"COUNT({a})", f"SUM({a}e)"]
            names += [f"min({col.name})", f"max({col.name})", f"non_null({col.name})",
                      f"date_sum({col.name})"]
    return inner, outer, names


def profile_sql(side: Side, spec: TableSpec, level: int, is_src: bool):
    inner, outer, names = _metric_exprs(side, spec, level, is_src)
    if spec.grain:
        inner = [f"{_part_expr(side, spec.grain)} AS p"] + inner
        outer = ["p"] + outer
        group = " GROUP BY p"
    else:
        group = ""
    inner_sql = ", ".join(inner) if inner else "1 AS one"
    sql = (f"SELECT {', '.join(outer)} FROM (SELECT {inner_sql} FROM {side.table}"
           f"{_where(side)}) x{group}")
    return sql, names


def bucket_sql(side: Side, spec: TableSpec, nb: int, pfilter, is_src: bool):
    inner, outer, names = _metric_exprs(side, spec, 2, is_src)
    b = side.d.bucket_expr(side.keys[0], spec.key_type, nb)
    sql = (f"SELECT b, {', '.join(outer)} FROM (SELECT {b} AS b, {', '.join(inner)} "
           f"FROM {side.table}{_where(side, [pfilter])}) x GROUP BY b")
    return sql, names


def fetch_sql(side: Side, spec: TableSpec, nb: int, buckets, pfilter, is_src: bool):
    d = side.d
    sel = [f"{k} AS k{i}" for i, k in enumerate(side.keys)]
    for i, col in enumerate(spec.columns):
        c = col.src if is_src else col.tgt
        sel.append(f"{d.to_day(c) if col.type == 'date' else c} AS c{i}")
    if nb <= 1 or buckets is None:
        return f"SELECT {', '.join(sel)} FROM {side.table}{_where(side, [pfilter])}"
    b = d.bucket_expr(side.keys[0], spec.key_type, nb)
    outer = [f"k{i}" for i in range(len(side.keys))] + [f"c{i}" for i in range(len(spec.columns))]
    return (f"SELECT {', '.join(outer)} FROM (SELECT {b} AS b, {', '.join(sel)} "
            f"FROM {side.table}{_where(side, [pfilter])}) x "
            f"WHERE b IN ({', '.join(str(x) for x in sorted(buckets))})")


# =============================================================================
# Normalisation helpers (make values from different engines comparable)
# =============================================================================
def _norm_scalar(v):
    if v is None:
        return None
    if isinstance(v, Decimal):
        v = float(v)
    if isinstance(v, float):
        if math.isnan(v):
            return None
        return int(v) if v.is_integer() else round(v, 8)
    if isinstance(v, (dt.datetime, dt.date)):
        return v.isoformat(sep=" ") if isinstance(v, dt.datetime) else v.isoformat()
    if isinstance(v, str):
        v = v.strip()
        return v or None
    return v


def _norm_date_str(v, with_time):
    if v is None:
        return None
    try:
        ts = pd.Timestamp(v)
    except (ValueError, TypeError):
        return str(v)
    return ts.strftime("%Y-%m-%d %H:%M:%S" if with_time else "%Y-%m-%d")


def _metric_equal(name, a, b, spec: TableSpec, rows):
    if a is None and b is None:
        return True
    if name.startswith(("min(", "max(")):
        col = next(c for c in spec.columns if name[4:-1] == c.name)
        return _norm_date_str(a, col.type == "datetime") == _norm_date_str(b, col.type == "datetime")
    try:
        fa, fb = float(a or 0), float(b or 0)
    except (TypeError, ValueError):
        return _norm_scalar(a) == _norm_scalar(b)
    if name.startswith("sum("):
        # sums are of values already rounded to `scale` in the DB, so they should agree
        # exactly; allow half a unit of the last decimal (+ float noise) plus tolerance.
        col = next(c for c in spec.columns if name[4:-1] == c.name)
        tol = 0.5 * 10 ** -col.scale + 1e-14 * max(abs(fa), abs(fb)) \
            + col.tolerance * max(rows or 1, 1)
    else:
        tol = 1e-9  # counts, lengths, signatures: exact
    return abs(fa - fb) <= tol


def _to_metric_map(rows, names, grain):
    out = {}
    for r in rows:
        if grain:
            pk, raw, vals = _norm_scalar(r[0]), r[0], r[1:]
            pk = None if pk is None else str(pk)
        else:
            pk, raw, vals = "ALL", None, r
        out[pk] = (raw, dict(zip(names, vals)))
    return out


def _norm_frame(df: pd.DataFrame, spec: TableSpec) -> pd.DataFrame:
    out = pd.DataFrame(index=df.index)
    for i, k in enumerate(spec.key_names):
        s = df[f"k{i}"]
        if spec.key_type == "number":
            out[k] = pd.to_numeric(s.map(lambda v: float(v) if isinstance(v, Decimal) else v),
                                   errors="coerce").map(
                lambda v: None if pd.isna(v) else (str(int(v)) if float(v).is_integer() else str(v)))
        else:
            out[k] = s.map(lambda v: None if v is None else str(v).strip())
    for i, col in enumerate(spec.columns):
        s = df[f"c{i}"]
        if col.type == "number":
            out[col.name] = pd.to_numeric(
                s.map(lambda v: float(v) if isinstance(v, Decimal) else v),
                errors="coerce").round(col.scale)
        elif col.type in ("date", "datetime"):
            out[col.name] = s.map(lambda v: _norm_date_str(v, col.type == "datetime"))
        else:
            def _t(v, upper=col.ignore_case):
                if v is None or (isinstance(v, float) and math.isnan(v)):
                    return None
                v = str(v).strip()
                return (v.upper() if upper else v) or None
            out[col.name] = s.map(_t).astype(object)
    return out


# =============================================================================
# Results
# =============================================================================
@dataclass
class PartitionResult:
    partition: str
    src_rows: int
    tgt_rows: int
    status: str
    failed_checks: list = field(default_factory=list)
    note: str = ""


@dataclass
class TableResult:
    name: str
    source: str
    target: str
    status: str = "MATCHED"
    src_rows: int = 0
    tgt_rows: int = 0
    partitions: list = field(default_factory=list)
    diffs: list = field(default_factory=list)
    diff_counts: dict = field(default_factory=dict)
    rows_fetched: int = 0
    spot_checked_buckets: int = 0
    seconds: float = 0.0
    error: str = ""
    notes: list = field(default_factory=list)


# =============================================================================
# Reconciliation engine
# =============================================================================
class Reconciler:
    def __init__(self, spec: TableSpec, settings: dict, level: int):
        self.spec = spec
        self.s = settings
        self.level = level
        self.bucket_size = int(settings.get("bucket_size", 20000))
        self.max_buckets = int(settings.get("max_buckets", 20000))
        self.max_diff_rows = int(settings.get("max_diff_rows", 1000))
        self.max_fetch_rows = int(settings.get("max_fetch_rows", 2_000_000))
        self.spot = int(settings.get("spot_check_buckets", 2))
        self.fetch_batch = int(settings.get("fetch_batch_buckets", 25))
        self.res = TableResult(name=spec.name,
                               source=f"{spec.src.db.name}:{spec.src.table}",
                               target=f"{spec.tgt.db.name}:{spec.tgt.table}")
        self._pool = cf.ThreadPoolExecutor(max_workers=2)

    def _both(self, sql_src, sql_tgt):
        fs = self._pool.submit(self.spec.src.db.query, sql_src)
        ft = self._pool.submit(self.spec.tgt.db.query, sql_tgt)
        return fs.result(), ft.result()

    # ---- main -------------------------------------------------------------
    def run(self) -> TableResult:
        t0 = time.time()
        try:
            self._run()
        except Exception as e:  # report and carry on with the other tables
            log.exception("Table %s failed", self.spec.name)
            self.res.status, self.res.error = "ERROR", f"{type(e).__name__}: {e}"
        finally:
            self._pool.shutdown(wait=False)
        self.res.seconds = round(time.time() - t0, 2)
        return self.res

    def _run(self):
        spec = self.spec
        sq, names = profile_sql(spec.src, spec, self.level, True)
        tq, _ = profile_sql(spec.tgt, spec, self.level, False)
        log.info("[%s] level %d profile ...", spec.name, self.level)
        srows, trows = self._both(sq, tq)
        smap = _to_metric_map(srows, names, spec.grain)
        tmap = _to_metric_map(trows, names, spec.grain)

        mismatched = []
        self.praw = {}
        for pk in sorted(set(smap) | set(tmap), key=lambda x: (x is None, x)):
            sraw, sm = smap.get(pk, (None, {}))
            traw, tm = tmap.get(pk, (None, {}))
            self.praw[pk] = sraw if sraw is not None else traw
            sc, tc = int(sm.get("rows") or 0), int(tm.get("rows") or 0)
            self.res.src_rows += sc
            self.res.tgt_rows += tc
            failed = [n for n in names
                      if not _metric_equal(n, sm.get(n), tm.get(n), spec, max(sc, tc))]
            if "distinct_keys" in tm and tm.get("distinct_keys") is not None and \
                    int(tm["distinct_keys"]) < tc and "duplicate_keys_in_target" not in failed:
                failed.append("duplicate_keys_in_target")
            pr = PartitionResult(partition="NULL" if pk is None else pk,
                                 src_rows=sc, tgt_rows=tc,
                                 status="MISMATCH" if failed else "MATCHED",
                                 failed_checks=failed)
            if sc == 0 and tc > 0:
                pr.note = "Partition exists only in target"
            elif tc == 0 and sc > 0:
                pr.note = "Partition missing in target"
            self.res.partitions.append(pr)
            if failed:
                mismatched.append((pk, sraw if sraw is not None else traw, pr))

        if mismatched:
            self.res.status = "MISMATCH"
        log.info("[%s] %d partitions, %d mismatched", spec.name,
                 len(self.res.partitions), len(mismatched))

        if self.level >= 3:
            for pk, raw, pr in mismatched:
                if self.res.rows_fetched >= self.max_fetch_rows:
                    pr.note = (pr.note + "; " if pr.note else "") + \
                        "drill-down skipped (max_fetch_rows reached)"
                    continue
                self._drill(pk, raw, pr)
            self._spot_check({pk for pk, _, _ in mismatched})
            self._merge_moved_rows()

    def _merge_moved_rows(self):
        """A row whose partition column changed shows up as MISSING in one partition and
        EXTRA in another. Report it once, as the value change it really is."""
        if not self.spec.grain:
            return
        miss = {d["key"]: d for d in self.res.diffs if d["issue"] == "MISSING_IN_TARGET"}
        extra = {d["key"]: d for d in self.res.diffs if d["issue"] == "EXTRA_IN_TARGET"}
        moved = set(miss) & set(extra)
        if not moved:
            return
        drop = {id(miss[k]) for k in moved} | {id(extra[k]) for k in moved}
        self.res.diffs = [d for d in self.res.diffs if id(d) not in drop]
        for k in sorted(moved):
            self.res.diffs.append({"table": self.spec.name, "partition": extra[k]["partition"],
                                   "key": k, "issue": "MOVED_PARTITION",
                                   "column": self.spec.src.part_col,
                                   "source_value": miss[k]["partition"],
                                   "target_value": extra[k]["partition"]})
        c = self.res.diff_counts
        c["MISSING_IN_TARGET"] -= len(moved)
        c["EXTRA_IN_TARGET"] -= len(moved)
        c["MOVED_PARTITION"] = c.get("MOVED_PARTITION", 0) + len(moved)
        for k in ("MISSING_IN_TARGET", "EXTRA_IN_TARGET"):
            if not c[k]:
                del c[k]

    # ---- level 3: drill down ---------------------------------------------
    def _filters(self, pk, raw):
        spec = self.spec
        return (partition_filter(spec.src, spec.grain, pk, raw),
                partition_filter(spec.tgt, spec.grain, pk, raw))

    def _drill(self, pk, raw, pr: PartitionResult):
        spec = self.spec
        fs, ft = self._filters(pk, raw)
        biggest = max(pr.src_rows, pr.tgt_rows)
        nb = min(self.max_buckets, max(1, math.ceil(biggest / self.bucket_size)))
        if nb == 1:
            found = self._fetch_and_diff(pk, raw, 1, None)
        else:
            sq, names = bucket_sql(spec.src, spec, nb, fs, True)
            tq, _ = bucket_sql(spec.tgt, spec, nb, ft, False)
            srows, trows = self._both(sq, tq)
            sb = {int(r[0]): dict(zip(names, r[1:])) for r in srows}
            tb = {int(r[0]): dict(zip(names, r[1:])) for r in trows}
            bad = []
            for b in set(sb) | set(tb):
                s, t = sb.get(b, {}), tb.get(b, {})
                rows = max(int(s.get("rows") or 0), int(t.get("rows") or 0))
                if any(not _metric_equal(n, s.get(n), t.get(n), spec, rows) for n in names):
                    bad.append(b)
                elif len(spec.tgt.keys) == 1 and t.get("distinct_keys") is not None and \
                        int(t["distinct_keys"]) < int(t.get("rows") or 0):
                    bad.append(b)
            log.info("[%s] partition %s: %d/%d buckets differ", spec.name, pk, len(bad), nb)
            found = 0
            for i in range(0, len(bad), self.fetch_batch):
                if self.res.rows_fetched >= self.max_fetch_rows:
                    pr.note = (pr.note + "; " if pr.note else "") + "drill-down truncated"
                    break
                found += self._fetch_and_diff(pk, raw, nb, bad[i:i + self.fetch_batch])
        if found == 0:
            pr.note = (pr.note + "; " if pr.note else "") + \
                "Aggregates differ but no row-level difference found (check rounding/tolerance)"

    def _spot_check(self, skip: set):
        """Row-compare a few random buckets every run, to catch changes aggregates can't see."""
        if self.spot <= 0:
            return
        cands = [(p.partition, p) for p in self.res.partitions
                 if p.status == "MATCHED" and p.src_rows > 0
                 and (None if p.partition == "NULL" else p.partition) not in skip]
        if not cands:
            return
        rng = random.Random()
        for _ in range(self.spot):
            pk_str, pr = rng.choice(cands)
            pk = None if pk_str == "NULL" else pk_str
            raw = self.praw.get(pk, pk)
            nb = min(self.max_buckets, max(1, math.ceil(pr.src_rows / self.bucket_size)))
            b = rng.randrange(nb)
            found = self._fetch_and_diff(pk, raw, nb, [b] if nb > 1 else None)
            self.res.spot_checked_buckets += 1
            if found:
                pr.status = "MISMATCH"
                pr.failed_checks.append("spot_check_rows")
                pr.note = "Difference found by random row-level spot check"
                self.res.status = "MISMATCH"

    def _fetch_and_diff(self, pk, raw, nb, buckets) -> int:
        spec = self.spec
        fs, ft = self._filters(pk, raw)
        sq = fetch_sql(spec.src, spec, nb, buckets, fs, True)
        tq = fetch_sql(spec.tgt, spec, nb, buckets, ft, False)
        srows, trows = self._both(sq, tq)
        self.res.rows_fetched += len(srows) + len(trows)
        cols = [f"k{i}" for i in range(len(spec.key_names))] + \
               [f"c{i}" for i in range(len(spec.columns))]
        sdf = _norm_frame(pd.DataFrame(srows, columns=cols, dtype=object), spec)
        tdf = _norm_frame(pd.DataFrame(trows, columns=cols, dtype=object), spec)
        return self._diff(sdf, tdf, "NULL" if pk is None else pk)

    def _record(self, issue, part, key, column=None, sv=None, tv=None):
        self.res.diff_counts[issue] = self.res.diff_counts.get(issue, 0) + 1
        if len(self.res.diffs) < self.max_diff_rows:
            self.res.diffs.append({"table": self.spec.name, "partition": part,
                                   "key": key, "issue": issue, "column": column,
                                   "source_value": sv, "target_value": tv})

    def _diff(self, sdf, tdf, part) -> int:
        keys = self.spec.key_names
        keystr = lambda r: " | ".join(str(r[k]) for k in keys)  # noqa: E731
        before = sum(self.res.diff_counts.values())

        for label, df in (("DUPLICATE_KEY_IN_SOURCE", sdf), ("DUPLICATE_KEY_IN_TARGET", tdf)):
            dup = df[df.duplicated(keys, keep="first")]
            for _, r in dup.iterrows():
                self._record(label, part, keystr(r))
        sdf = sdf.drop_duplicates(keys)
        tdf = tdf.drop_duplicates(keys)

        m = sdf.merge(tdf, on=keys, how="outer", suffixes=("__s", "__t"), indicator=True)
        for _, r in m[m["_merge"] == "left_only"].iterrows():
            self._record("MISSING_IN_TARGET", part, keystr(r))
        for _, r in m[m["_merge"] == "right_only"].iterrows():
            self._record("EXTRA_IN_TARGET", part, keystr(r))

        both = m[m["_merge"] == "both"]
        for col in self.spec.columns:
            s, t = both[f"{col.name}__s"], both[f"{col.name}__t"]
            sn, tn = s.isna(), t.isna()
            if col.type == "number":
                diff = (sn != tn) | (~sn & ~tn & ((s.astype(float) - t.astype(float)).abs()
                                                  > col.tolerance + 1e-9))
            else:
                diff = (sn != tn) | (~sn & ~tn & (s.astype(object) != t.astype(object)))
            for idx in both.index[diff.to_numpy()]:
                r = both.loc[idx]
                sv, tv = r[f"{col.name}__s"], r[f"{col.name}__t"]
                self._record("VALUE_MISMATCH", part, keystr(r), col.name,
                             None if pd.isna(sv) else sv, None if pd.isna(tv) else tv)
        return sum(self.res.diff_counts.values()) - before


# =============================================================================
# Reports
# =============================================================================
def build_frames(results: list[TableResult], run_id: str, run_ts: str):
    summary = pd.DataFrame([{
        "run_id": run_id, "run_time": run_ts, "table": r.name, "status": r.status,
        "source": r.source, "target": r.target, "source_rows": r.src_rows,
        "target_rows": r.tgt_rows, "row_diff": r.tgt_rows - r.src_rows,
        "partitions": len(r.partitions),
        "partitions_mismatched": sum(p.status != "MATCHED" for p in r.partitions),
        "missing_in_target": r.diff_counts.get("MISSING_IN_TARGET", 0),
        "extra_in_target": r.diff_counts.get("EXTRA_IN_TARGET", 0),
        "value_mismatches": r.diff_counts.get("VALUE_MISMATCH", 0),
        "moved_partition": r.diff_counts.get("MOVED_PARTITION", 0),
        "duplicate_keys": r.diff_counts.get("DUPLICATE_KEY_IN_TARGET", 0)
                          + r.diff_counts.get("DUPLICATE_KEY_IN_SOURCE", 0),
        "rows_fetched_for_drilldown": r.rows_fetched,
        "spot_checked_buckets": r.spot_checked_buckets,
        "seconds": r.seconds, "error": r.error} for r in results])
    parts = pd.DataFrame([{
        "run_id": run_id, "table": r.name, "partition": p.partition,
        "source_rows": p.src_rows, "target_rows": p.tgt_rows,
        "row_diff": p.tgt_rows - p.src_rows, "status": p.status,
        "failed_checks": ", ".join(p.failed_checks), "note": p.note}
        for r in results for p in r.partitions],
        columns=["run_id", "table", "partition", "source_rows", "target_rows", "row_diff",
                 "status", "failed_checks", "note"])
    diffs = pd.DataFrame([{"run_id": run_id, **d} for r in results for d in r.diffs],
                         columns=["run_id", "table", "partition", "key", "issue", "column",
                                  "source_value", "target_value"])
    return summary, parts, diffs


DB_COLUMN_NAMES = {   # avoid SQL reserved words (TABLE, KEY, PARTITION, COLUMN ...)
    "table": "TABLE_NAME", "source": "SOURCE_TABLE", "target": "TARGET_TABLE",
    "partition": "PARTITION_KEY", "key": "RECORD_KEY", "column": "COLUMN_NAME",
    "error": "ERROR_MESSAGE", "seconds": "DURATION_SECONDS",
    "rows_fetched_for_drilldown": "ROWS_FETCHED"}


def db_frame(df: pd.DataFrame, run_ts: str, as_text=False) -> pd.DataFrame:
    """Report frame -> database-ready frame: safe UPPERCASE names, real NULLs, typed run time."""
    out = df.copy()
    if "run_time" not in out.columns:
        out.insert(1, "run_time", run_ts)
    out["run_time"] = pd.to_datetime(out["run_time"])
    if as_text:
        for c in ("source_value", "target_value"):
            out[c] = out[c].map(lambda v: None if v is None or (isinstance(v, float) and
                                math.isnan(v)) else str(v)[:4000])
    for c in out.columns:
        if out[c].dtype == object:
            out[c] = out[c].map(lambda v: v[:4000] if isinstance(v, str) else v)
            out[c] = out[c].where(out[c].notna() & (out[c] != ""), None)
    out.columns = [DB_COLUMN_NAMES.get(c, c).upper() for c in out.columns]
    return out


def write_json(path, run_id, run_ts, level, summary, parts, diffs):
    """Machine-readable copy of the run, used by the web app."""
    import json

    def recs(df):
        return df.astype(object).where(df.notna(), None).to_dict(orient="records")
    doc = {"run_id": run_id, "run_time": run_ts, "level": level,
           "summary": recs(summary), "partitions": recs(parts), "differences": recs(diffs)}
    tmp = path + ".tmp"
    with open(tmp, "w", encoding="utf-8") as f:
        json.dump(doc, f, default=str)
    os.replace(tmp, path)


def write_excel(path, summary, parts, diffs):
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter
    fills = {"MATCHED": "C6EFCE", "MISMATCH": "FFC7CE", "ERROR": "FFEB9C"}
    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        summary.to_excel(xw, sheet_name="Summary", index=False)
        parts[parts["status"] != "MATCHED"].to_excel(xw, sheet_name="Mismatched partitions", index=False)
        diffs.astype(str).replace({"None": "", "nan": ""}).to_excel(xw, sheet_name="Differences", index=False)
        parts.to_excel(xw, sheet_name="All partitions", index=False)
        for ws in xw.book.worksheets:
            ws.freeze_panes = "A2"
            for c in ws[1]:
                c.font = Font(bold=True, color="FFFFFF")
                c.fill = PatternFill("solid", fgColor="1F4E78")
            heads = [c.value for c in ws[1]]
            if "status" in heads:
                si = heads.index("status") + 1
                for row in ws.iter_rows(min_row=2):
                    f = fills.get(row[si - 1].value)
                    if f:
                        row[si - 1].fill = PatternFill("solid", fgColor=f)
            for row in ws.iter_rows(min_row=2):
                for c in row:
                    if c.data_type == "f":   # never let data become a live Excel formula
                        c.data_type = "s"
            for i, col in enumerate(ws.columns, 1):
                width = max((len(str(c.value)) for c in list(col)[:500] if c.value is not None),
                            default=8)
                ws.column_dimensions[get_column_letter(i)].width = min(max(width + 2, 10), 60)


def write_html(path, summary, parts, diffs, run_ts, max_rows=50):
    e = html.escape
    total = len(summary)
    ok = int((summary["status"] == "MATCHED").sum()) if total else 0
    bad = int((summary["status"] == "MISMATCH").sum()) if total else 0
    err = int((summary["status"] == "ERROR").sum()) if total else 0

    def chip(s):
        return f'<span class="chip {s.lower()}">{e(s)}</span>'

    def table(df, cols, status_col="status"):
        if df.empty:
            return '<p class="muted">None.</p>'
        h = "".join(f"<th>{e(c.replace('_', ' '))}</th>" for c in cols)
        body = []
        for _, r in df.iterrows():
            tds = []
            for c in cols:
                v = r[c]
                v = "" if v is None or (isinstance(v, float) and math.isnan(v)) else v
                if c == status_col:
                    tds.append(f"<td>{chip(str(v))}</td>")
                elif isinstance(v, (int, float)) and not isinstance(v, bool):
                    tds.append(f'<td class="num">{v:,}</td>' if isinstance(v, int)
                               else f'<td class="num">{v:,.2f}</td>')
                else:
                    tds.append(f"<td>{e(str(v))}</td>")
            body.append("<tr>" + "".join(tds) + "</tr>")
        return f'<div class="wrap"><table><thead><tr>{h}</tr></thead><tbody>{"".join(body)}</tbody></table></div>'

    sections = []
    for _, s in summary.iterrows():
        if s["status"] == "MATCHED":
            continue
        tp = parts[(parts["table"] == s["table"]) & (parts["status"] != "MATCHED")]
        td = diffs[diffs["table"] == s["table"]].head(max_rows)
        extra = (f'<p class="err">{e(s["error"])}</p>' if s["error"] else "")
        sections.append(
            f'<section><h2>{e(s["table"])} {chip(s["status"])}</h2>{extra}'
            f'<h3>Mismatched partitions</h3>'
            f'{table(tp, ["partition", "source_rows", "target_rows", "row_diff", "failed_checks", "note"])}'
            f'<h3>Differences (first {max_rows})</h3>'
            f'{table(td, ["partition", "key", "issue", "column", "source_value", "target_value"], "_")}'
            f'</section>')

    doc = f"""<!doctype html><html><head><meta charset="utf-8">
<meta name="viewport" content="width=device-width,initial-scale=1">
<title>Reconciliation Report</title><style>
:root{{--bg:#f6f7f9;--card:#fff;--ink:#1d2330;--muted:#667085;--line:#e4e7ec;
--ok:#067647;--okbg:#dcfae6;--bad:#b42318;--badbg:#fee4e2;--warn:#b54708;--warnbg:#fef0c7}}
body{{margin:0;background:var(--bg);color:var(--ink);font:14px/1.5 system-ui,Segoe UI,Arial,sans-serif}}
main{{max-width:1200px;margin:0 auto;padding:24px 16px}}
h1{{margin:0 0 4px;font-size:22px}} h2{{font-size:17px;margin:0 0 8px}} h3{{font-size:13px;color:var(--muted);text-transform:uppercase;letter-spacing:.04em;margin:16px 0 6px}}
.muted{{color:var(--muted)}} .err{{color:var(--bad)}}
.kpis{{display:grid;grid-template-columns:repeat(auto-fit,minmax(150px,1fr));gap:12px;margin:18px 0}}
.kpi,section{{background:var(--card);border:1px solid var(--line);border-radius:10px;padding:14px 16px}}
section{{margin-top:16px}} .kpi b{{display:block;font-size:26px}}
.wrap{{overflow-x:auto}} table{{border-collapse:collapse;width:100%;font-size:13px}}
th,td{{text-align:left;padding:6px 10px;border-bottom:1px solid var(--line);white-space:nowrap}}
th{{color:var(--muted);font-weight:600;background:#fafbfc}} td.num{{text-align:right;font-variant-numeric:tabular-nums}}
.chip{{display:inline-block;padding:1px 8px;border-radius:99px;font-size:12px;font-weight:600}}
.matched{{color:var(--ok);background:var(--okbg)}} .mismatch{{color:var(--bad);background:var(--badbg)}} .error{{color:var(--warn);background:var(--warnbg)}}
</style></head><body><main>
<h1>Data Reconciliation Report</h1><div class="muted">Run at {e(run_ts)}</div>
<div class="kpis"><div class="kpi"><span class="muted">Tables checked</span><b>{total}</b></div>
<div class="kpi"><span class="muted">Matched</span><b style="color:var(--ok)">{ok}</b></div>
<div class="kpi"><span class="muted">Mismatched</span><b style="color:var(--bad)">{bad}</b></div>
<div class="kpi"><span class="muted">Errors</span><b style="color:var(--warn)">{err}</b></div></div>
<section><h2>Summary</h2>{table(summary, ["table", "status", "source_rows", "target_rows", "row_diff",
        "partitions_mismatched", "missing_in_target", "extra_in_target", "value_mismatches",
        "moved_partition", "duplicate_keys", "seconds"])}</section>
{"".join(sections)}
</main></body></html>"""
    with open(path, "w", encoding="utf-8") as f:
        f.write(doc)


# =============================================================================
# CLI
# =============================================================================
def connection_hint(err: Exception, cfg: dict) -> str:
    """Turn common driver errors into a plain-language fix."""
    m = f"{type(err).__name__}: {err}"
    low = m.lower()
    if isinstance(err, ValueError):
        return "Edit recon_config.yaml / .env as the error says, then run --check again."
    host = cfg.get("host") or cfg.get("server") or urllib.parse.urlsplit(
        cfg.get("url", "").split("?")[0].replace("{password}", "x")).hostname
    if "getaddrinfo" in low or "name or service not known" in low or "11001" in low \
            or "could not translate host" in low or "dpy-6005" in low and "resolve" in low:
        return (f"Server name '{host}' not found on the network. Put the real host name or IP "
                f"in recon_config.yaml, and check "
                f"you are on the office network/VPN.")
    if "im002" in low or "data source name not found" in low:
        try:
            import pyodbc
            drv = pyodbc.drivers()
        except Exception:
            drv = []
        return ("The ODBC driver named in the url is not installed. Installed drivers on this PC: "
                f"{drv or 'NONE - install Microsoft ODBC Driver 18 for SQL Server'}. Put one of "
                "them as `driver:` in recon_config.yaml, copied exactly.")
    if "ora-01017" in low or "login failed" in low or "18456" in low \
            or "password authentication failed" in low or "access denied" in low:
        return ("Server reached, but user name or password was rejected. Check the user in the "
                f"url and the real password for {cfg.get('password_env')} in .env.")
    if "dpy-3015" in low or "dpy-3001" in low or "dpy-4011" in low or "dpy-3010" in low:
        return ("This Oracle database needs 'thick mode'. Install Oracle Instant Client and add "
                "`thick_mode: true` (and `client_lib_dir: C:\\oracle\\instantclient_xx`) to the "
                "dwh connection in recon_config.yaml.")
    if "ora-12514" in low or "dpy-6001" in low or "service" in low and "not" in low and "ora-" in low:
        return "Oracle listener reached, but the service_name in the url is wrong. Check tnsnames.ora."
    if "timed out" in low or "timeout" in low or "10060" in low or "refused" in low \
            or "hyt00" in low or "08001" in low:
        return ("Server name found but no answer on that port. Check the port number, that the "
                "database is running, and that the firewall allows this PC to connect.")
    if "cannot open database" in low or "4060" in low:
        return "Logged in, but the database name in the url is wrong or this user has no access to it."
    if "password not found" in low:
        return m.split(": ", 1)[1]
    if "no module named" in low:
        return f"Python driver missing: {m.split('No module named')[-1].strip()} - pip install it."
    return "Unrecognised error, see the message above."


def check_connections(conns: dict) -> int:
    print("\nChecking connections ...\n")
    bad = 0
    for name, c in conns.items():
        t0 = time.time()
        try:
            db = Database(name, c)
            db.query("SELECT 1 FROM DUAL" if db.type == "oracle" else "SELECT 1")
            print(f"  OK    {name:12} ({c.get('type')})  {time.time() - t0:.1f}s")
        except Exception as e:
            bad += 1
            msg = str(e).splitlines()[0][:300]
            print(f"  FAIL  {name:12} ({c.get('type')})\n        error: {type(e).__name__}: {msg}"
                  f"\n        fix  : {connection_hint(e, c)}\n")
    print("\nAll connections OK." if not bad else f"\n{bad} connection(s) failed.")
    return 2 if bad else 0


def load_env_file(path) -> int | None:
    """Read KEY=VALUE lines into the environment (no extra package needed).
    Values already set in the environment win. Returns number of keys read, or None."""
    if not path or not os.path.isfile(path):
        return None
    n = 0
    with open(path, encoding="utf-8-sig") as f:      # utf-8-sig: tolerate Notepad's BOM
        for line in f:
            line = line.strip()
            if not line or line.startswith("#") or "=" not in line:
                continue
            if line.startswith("export "):
                line = line[7:]
            k, v = line.split("=", 1)
            k, v = k.strip(), v.strip()
            if len(v) >= 2 and v[0] == v[-1] and v[0] in "\"'":
                v = v[1:-1]
            os.environ.setdefault(k, v)
            n += 1
    return n


def main(argv=None):
    ap = argparse.ArgumentParser(description="Reconcile DWH source tables with destination tables")
    ap.add_argument("-c", "--config", required=True, help="YAML config file")
    ap.add_argument("-t", "--table", action="append", help="Only run these table names (repeatable)")
    ap.add_argument("-l", "--level", type=int, choices=(1, 2, 3), default=None,
                    help="1=counts, 2=+aggregates, 3=+row-level drill-down (default from config, else 3)")
    ap.add_argument("--show-sql", action="store_true", help="Print the generated SQL and exit")
    ap.add_argument("--check", action="store_true",
                    help="Test every connection (host, driver, login) and exit")
    ap.add_argument("--run-id", help="Use this run id (the web app sets it); default: timestamp")
    ap.add_argument("--env-file", help="File with passwords (default: .env next to the config, "
                                       "then .env in the current folder)")
    ap.add_argument("-v", "--verbose", action="store_true")
    a = ap.parse_args(argv)

    logging.basicConfig(level=logging.DEBUG if a.verbose else logging.INFO,
                        format="%(asctime)s %(levelname)s %(message)s", datefmt="%H:%M:%S")
    env_candidates = [a.env_file] if a.env_file else [
        os.path.join(os.path.dirname(os.path.abspath(a.config)), ".env"), ".env"]
    for p in env_candidates:
        n = load_env_file(p)
        if n is not None:
            log.info("Loaded %d value(s) from %s", n, os.path.abspath(p))
            break
    else:
        if a.env_file:
            log.error("Env file not found: %s", a.env_file)
            return 2

    a.config = os.path.abspath(a.config)
    # Relative paths in the config (output_dir, sqlite files) are relative to the config's folder
    os.chdir(os.path.dirname(a.config))
    with open(a.config, encoding="utf-8") as f:
        cfg = yaml.safe_load(f)
    settings = cfg.get("settings") or {}
    level = a.level or int(settings.get("level", 3))

    out = settings.get("output_dir", "recon_output")
    os.makedirs(out, exist_ok=True)
    fh = logging.FileHandler(os.path.join(out, "recon.log"), encoding="utf-8")
    fh.setFormatter(logging.Formatter("%(asctime)s %(levelname)s %(message)s"))
    logging.getLogger().addHandler(fh)
    log.info("=== Run started: config=%s level=%d ===", a.config, level)

    if a.check:
        return check_connections(cfg["connections"])

    try:
        dbs = {n: Database(n, c) for n, c in cfg["connections"].items()}
        tables = [t for t in cfg["tables"] if t.get("enabled", True)
                  and (not a.table or t["name"] in a.table)]
        if not tables:
            log.error("No tables selected")
            return 2
        specs = [build_spec(t, dbs, settings) for t in tables]
    except (ValueError, KeyError, ImportError, FileNotFoundError) as e:
        log.error("Setup problem: %s", e if not isinstance(e, KeyError)
                  else f"missing setting {e} in recon_config.yaml")
        log.error("Fix it, then test connections with:  python recon.py -c %s --check", a.config)
        return 2

    if a.show_sql:
        for sp in specs:
            print(f"\n-- ===== {sp.name} =====")
            for label, side, is_src in (("SOURCE", sp.src, True), ("TARGET", sp.tgt, False)):
                print(f"-- {label} profile ({side.db.type})")
                print(profile_sql(side, sp, level, is_src)[0] + ";")
                print(f"-- {label} bucket fingerprint (example: 100 buckets)")
                print(bucket_sql(side, sp, 100, None, is_src)[0] + ";")
        return 0

    run_ts = dt.datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    run_id = a.run_id or dt.datetime.now().strftime("%Y%m%d_%H%M%S")
    t0 = time.time()
    with cf.ThreadPoolExecutor(max_workers=int(settings.get("parallel_tables", 4))) as ex:
        results = list(ex.map(lambda sp: Reconciler(sp, settings, level).run(), specs))
    log.info("Finished %d tables in %.1fs", len(results), time.time() - t0)

    summary, parts, diffs = build_frames(results, run_id, run_ts)
    xlsx = os.path.join(out, f"recon_{run_id}.xlsx")
    htm = os.path.join(out, f"recon_{run_id}.html")
    write_excel(xlsx, summary, parts, diffs)
    write_html(htm, summary, parts, diffs, run_ts)
    hist = os.path.join(out, "recon_history.csv")
    summary.to_csv(hist, mode="a", header=not os.path.exists(hist), index=False)
    write_json(os.path.join(out, f"recon_{run_id}.json"), run_id, run_ts, level, summary, parts, diffs)

    rcfg = settings.get("results_db")
    db_write_failed = False
    if rcfg:
        try:
            db = dbs[rcfg["connection"]]
            pre, schema = rcfg.get("table_prefix", "RECON_"), rcfg.get("schema")
            db.write_tables({f"{pre}SUMMARY": db_frame(summary, run_ts),
                             f"{pre}PARTITIONS": db_frame(parts, run_ts),
                             f"{pre}DIFFERENCES": db_frame(diffs, run_ts, as_text=True)}, schema)
            log.info("Results saved to database '%s': %sSUMMARY %d rows, %sPARTITIONS %d rows, "
                     "%sDIFFERENCES %d rows", db.name, pre, len(summary), pre, len(parts),
                     pre, len(diffs))
        except Exception as e:
            db_write_failed = True
            log.error("Could not write results to database (reports were still saved): %s", e)

    print("\n" + "=" * 78)
    for r in results:
        dc = ", ".join(f"{k}={v}" for k, v in sorted(r.diff_counts.items())) or "-"
        print(f"{r.status:9} {r.name:28} src={r.src_rows:>12,} tgt={r.tgt_rows:>12,} "
              f"{r.seconds:>7.1f}s  {dc}{('  ' + r.error) if r.error else ''}")
    print("=" * 78)
    print(f"Excel : {xlsx}\nHTML  : {htm}\nHistory: {hist}\nRUN_ID: {run_id}")

    if db_write_failed or any(r.status == "ERROR" for r in results):
        return 2
    return 1 if any(r.status == "MISMATCH" for r in results) else 0


if __name__ == "__main__":
    sys.exit(main())